#!/usr/bin/env python3
"""Generate Word review packets for non-technical review.

The generated files are review artifacts only. They are intended for comments and
tracked changes in Word; accepted edits should be applied back to the source
files such as utilities/descriptions.yaml and docs/*.tex.

Run from repository root or from docs/:

    python docs/generate_review_packet.py

Outputs are written to docs/review/ by default.
"""
from __future__ import annotations

import argparse
import re
import subprocess
import sys
from html import unescape
from pathlib import Path
from typing import Iterable

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
except ImportError as exc:  # pragma: no cover - user environment guard
    raise SystemExit(
        "python-docx is required to generate review packets. "
        "Install dependencies with `pip install -r requirements.txt`."
    ) from exc

ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = ROOT / "docs"
REVIEW_DIR = DOCS_DIR / "review"

if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from utilities.descriptions import (  # noqa: E402
    ANALYSIS_CONFIG_LABELS,
    ANALYSIS_PARAMETER_DESCRIPTIONS,
    APPROACH_DESCRIPTIONS,
    CROSS_VALIDATION_DESCRIPTION_HTML,
    CROSS_VALIDATION_SELECTION_GUIDANCE,
    CROSS_VALIDATION_SUMMARY_HTML,
    CV_SCHEME_DESCRIPTIONS,
    CV_SCHEME_LABELS,
    GRID_SEARCH_DESCRIPTION_HTML,
    METRIC_DESCRIPTIONS,
    MODAL_HELP,
    PARAMETER_DIAGNOSTICS_COMPARISON_HTML,
    PREPROCESSING_DESCRIPTIONS,
    RELATIVE_MODEL_EVIDENCE_TEXT,
    TOOL_CAPABILITY_SUMMARY,
    TOOL_PURPOSE_TEXT,
)

try:  # pragma: no cover - documentation fallback only
    from metrics import metrics
except Exception:
    metrics = {}

try:  # pragma: no cover - documentation fallback only
    from models import approaches, models
except Exception:
    approaches = {}
    models = {}


def html_to_text(value: str) -> str:
    """Convert short HTML/Jinja-free fragments to plain text."""
    if not value:
        return ""
    text = value.strip()
    text = re.sub(r"</p>\s*<p>", "\n\n", text)
    text = re.sub(r"<br\s*/?>", "\n", text)
    text = re.sub(r"<li>", "- ", text)
    text = re.sub(r"</li>", "\n", text)
    text = re.sub(r"</?(ul|ol)>", "\n", text)
    text = re.sub(r"<strong>(.*?)</strong>", r"\1", text, flags=re.S)
    text = re.sub(r"<em>(.*?)</em>", r"\1", text, flags=re.S)
    text = re.sub(r"<[^>]+>", "", text)
    text = unescape(text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def add_title(doc: Document, title: str, subtitle: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(18)
    if subtitle:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(subtitle)
        run.italic = True
        run.font.size = Pt(10)
    doc.add_paragraph()


def add_paragraphs(doc: Document, text: str) -> None:
    for paragraph in [p.strip() for p in re.split(r"\n\s*\n", text or "") if p.strip()]:
        if paragraph.startswith("- "):
            for item in paragraph.splitlines():
                item = item.strip()
                if item.startswith("- "):
                    doc.add_paragraph(item[2:].strip(), style="List Bullet")
                elif item:
                    doc.add_paragraph(item)
        else:
            doc.add_paragraph(paragraph)


def add_table(doc: Document, headers: list[str], rows: Iterable[Iterable[object]]) -> None:
    rows = [list(row) for row in rows]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = str(header)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        normalized = list(row)
        if len(normalized) < len(headers):
            normalized.extend([""] * (len(headers) - len(normalized)))
        elif len(normalized) > len(headers):
            normalized = normalized[:len(headers) - 1] + [" | ".join(str(v) for v in normalized[len(headers) - 1:])]
        cells = table.add_row().cells
        for idx, value in enumerate(normalized):
            cells[idx].text = html_to_text(str(value))
    doc.add_paragraph()


def direction_label(direction: str) -> str:
    if direction == "higher":
        return "Higher is better"
    if direction == "lower":
        return "Lower is better"
    return direction or ""


def model_approach_labels(model_key: str) -> str:
    labels: list[str] = []
    for approach_key, approach_info in approaches.items():
        if model_key in approach_info.get("models", []):
            labels.append(approach_info.get("display_name", approach_key))
    return ", ".join(labels)


def build_app_content_docx(path: Path) -> None:
    doc = Document()
    add_title(
        doc,
        "IVIVC App Content Review Packet",
        "Review artifact for comments and tracked changes; source files remain authoritative.",
    )

    doc.add_heading("How to use this review packet", level=1)
    add_paragraphs(
        doc,
        "This document collects user-facing text and compact summaries from the app source. "
        "Reviewers may use Word comments or tracked changes to suggest wording, clarity, or formatting changes. "
        "Accepted edits should be applied back to the app source files rather than treating this document as the source of truth."
    )
    doc.add_paragraph("Screenshots and live-app testing should be reviewed separately when layout or visual behavior is the concern.")

    doc.add_heading("Purpose and capability summary", level=1)
    add_paragraphs(doc, TOOL_PURPOSE_TEXT)
    for item in TOOL_CAPABILITY_SUMMARY:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Input page content", level=1)
    doc.add_heading("Data upload", level=2)
    add_paragraphs(
        doc,
        "Fitting dataset: required dataset containing in vitro and in vivo data for model fitting.\n\n"
        "Prediction dataset: optional dataset for making predictions using fitted models.\n\n"
        "Supported formats: CSV and Excel (.xlsx, .xls)."
    )

    doc.add_heading("Analysis approaches", level=2)
    approach_rows = []
    for approach_key, desc in APPROACH_DESCRIPTIONS.items():
        model_names = [models[m].get("display_name", m) for m in approaches.get(approach_key, {}).get("models", []) if m in models]
        approach_rows.append([
            desc.get("title", approach_key),
            html_to_text(desc.get("method_html", "")),
            ", ".join(model_names) if model_names else "No parametric model selection",
        ])
    add_table(doc, ["Approach", "User-facing method description", "Models"], approach_rows)

    doc.add_heading("Preprocessing options", level=2)
    for category, options in PREPROCESSING_DESCRIPTIONS.items():
        doc.add_heading(category.replace("_", " ").title(), level=3)
        add_table(
            doc,
            ["Option", "Description"],
            [[info.get("display_name", key), info.get("description", "")] for key, info in options.items()],
        )

    doc.add_heading("Advanced analysis settings", level=2)
    doc.add_heading("Cross-validation", level=3)
    add_paragraphs(doc, html_to_text(CROSS_VALIDATION_DESCRIPTION_HTML))
    guidance_rows = [
        [row.get("scheme", ""), row.get("use_when", ""), row.get("caution", "")]
        for row in CROSS_VALIDATION_SELECTION_GUIDANCE
    ]
    if guidance_rows:
        doc.add_heading("Choosing a cross-validation scheme", level=4)
        add_table(doc, ["Scheme", "Use when", "Main caution"], guidance_rows)
    add_table(
        doc,
        ["Scheme", "Internal key", "Description"],
        [[label, key, CV_SCHEME_DESCRIPTIONS.get(key, "")] for key, label in CV_SCHEME_LABELS.items()],
    )
    doc.add_heading("Grid-search initialization", level=3)
    add_paragraphs(doc, html_to_text(GRID_SEARCH_DESCRIPTION_HTML))
    analysis_rows = [[label, key, ANALYSIS_PARAMETER_DESCRIPTIONS.get(key, "")] for key, label in ANALYSIS_CONFIG_LABELS.items()]
    add_table(doc, ["Parameter", "Internal key", "Description"], analysis_rows)

    doc.add_heading("Performance metrics", level=2)
    metric_rows = []
    if metrics:
        for metric_key, metric_info in metrics.items():
            desc = METRIC_DESCRIPTIONS.get(metric_key, {}).get("description", "")
            metric_rows.append([
                metric_info.get("display_name", metric_key),
                metric_key,
                direction_label(metric_info.get("better_direction", "")),
                desc,
            ])
    else:
        for metric_key, desc in METRIC_DESCRIPTIONS.items():
            metric_rows.append([metric_key, metric_key, "", desc.get("description", "")])
    add_table(doc, ["Metric", "Internal key", "Preferred direction", "Description"], metric_rows)

    doc.add_heading("Model definitions", level=1)
    model_rows = []
    for model_key, model_info in sorted(models.items(), key=lambda item: item[1].get("display_name", item[0])):
        model_rows.append([
            model_info.get("display_name", model_key),
            model_approach_labels(model_key),
            model_info.get("latex_equation", ""),
            model_info.get("description", ""),
        ])
    add_table(doc, ["Model", "Used in", "Functional form", "Description"], model_rows)

    doc.add_heading("Results and report explanatory text", level=1)
    doc.add_heading("Relative model evidence", level=2)
    add_paragraphs(doc, RELATIVE_MODEL_EVIDENCE_TEXT)
    doc.add_heading("Cross-validation summary", level=2)
    add_paragraphs(doc, html_to_text(CROSS_VALIDATION_SUMMARY_HTML))
    doc.add_heading("Fitted parameter diagnostics comparison", level=2)
    add_paragraphs(doc, html_to_text(PARAMETER_DIAGNOSTICS_COMPARISON_HTML))

    doc.add_heading("Modal and help text", level=1)
    for page_key, page_content in MODAL_HELP.items():
        doc.add_heading(page_key.replace("_", " ").title(), level=2)
        for key, value in page_content.items():
            doc.add_heading(key.replace("_", " ").title(), level=3)
            add_paragraphs(doc, html_to_text(str(value)))

    doc.add_heading("Reviewer notes", level=1)
    doc.add_paragraph("Use this space for general comments that do not map cleanly to a specific section above.")
    for _ in range(6):
        doc.add_paragraph(" ")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


_LATEX_REPLACEMENTS = [
    (r"\\textsubscript\{([^{}]*)\}", r"\1"),
    (r"\\textbf\{([^{}]*)\}", r"\1"),
    (r"\\emph\{([^{}]*)\}", r"\1"),
    (r"\\cite\{([^{}]*)\}", lambda m: "[cite: " + m.group(1).replace(",", "; ") + "]"),
    (r"\\url\{([^{}]*)\}", r"\1"),
]


def latex_to_plain(text: str) -> str:
    text = text.strip()
    for pattern, repl in _LATEX_REPLACEMENTS:
        text = re.sub(pattern, repl, text)
    text = text.replace(r"\(", "").replace(r"\)", "")
    text = text.replace(r"\[", "").replace(r"\]", "")
    text = text.replace("$", "")
    text = text.replace(r"\&", "&")
    text = text.replace(r"\%", "%")
    text = text.replace(r"\_", "_")
    text = text.replace(r"\#", "#")
    text = text.replace(r"\pm", "±")
    text = text.replace(r"\Delta", "Delta")
    text = text.replace(r"\times", "x")
    text = text.replace(r"\le", "<=")
    text = text.replace(r"\ge", ">=")
    text = re.sub(r"\\[a-zA-Z]+", "", text)
    text = text.replace("{", "").replace("}", "")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def split_table_row(line: str) -> list[str] | None:
    if " & " not in line or not line.rstrip().endswith(r"\\"):
        return None
    line = line.rstrip()[:-2].strip()
    if not line or line.startswith("\\"):
        return None
    return [latex_to_plain(part.strip()) for part in line.split("&")]


def add_latex_table(doc: Document, block: list[str]) -> None:
    rows: list[list[str]] = []
    for line in block:
        stripped = line.strip()
        if any(token in stripped for token in [r"\toprule", r"\midrule", r"\bottomrule", r"\endfirsthead", r"\endhead"]):
            continue
        row = split_table_row(stripped)
        if row and row not in rows:
            rows.append(row)
    if not rows:
        return
    headers = rows[0]
    data_rows = rows[1:]
    add_table(doc, headers, data_rows)


def render_latex_file(doc: Document, path: Path, heading_prefix: str | None = None) -> None:
    if heading_prefix:
        doc.add_heading(heading_prefix, level=1)
    lines = path.read_text(encoding="utf-8").splitlines()
    paragraph_buffer: list[str] = []
    in_itemize = False
    in_enumerate = False
    table_block: list[str] | None = None

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if paragraph_buffer:
            text = latex_to_plain(" ".join(paragraph_buffer))
            if text:
                doc.add_paragraph(text)
            paragraph_buffer = []

    for line in lines:
        stripped = line.strip()
        if not stripped or stripped.startswith("%"):
            flush_paragraph()
            continue

        if table_block is not None:
            if stripped.startswith(r"\end{longtable}"):
                add_latex_table(doc, table_block)
                table_block = None
            else:
                table_block.append(stripped)
            continue
        if stripped.startswith(r"\begin{longtable}"):
            flush_paragraph()
            table_block = []
            continue

        heading_match = re.match(r"\\(section|subsection|subsubsection)\{(.+)\}", stripped)
        if heading_match:
            flush_paragraph()
            level = {"section": 1, "subsection": 2, "subsubsection": 3}[heading_match.group(1)]
            doc.add_heading(latex_to_plain(heading_match.group(2)), level=level)
            continue

        if stripped == r"\begin{itemize}":
            flush_paragraph()
            in_itemize = True
            continue
        if stripped == r"\end{itemize}":
            flush_paragraph()
            in_itemize = False
            continue
        if stripped == r"\begin{enumerate}":
            flush_paragraph()
            in_enumerate = True
            continue
        if stripped == r"\end{enumerate}":
            flush_paragraph()
            in_enumerate = False
            continue

        if stripped.startswith(r"\item"):
            flush_paragraph()
            item_text = latex_to_plain(stripped.replace(r"\item", "", 1))
            style = "List Bullet" if in_itemize else "List Number" if in_enumerate else None
            doc.add_paragraph(item_text, style=style)
            continue

        paragraph_buffer.append(stripped)

    flush_paragraph()


def parse_bib_entries(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    content = path.read_text(encoding="utf-8")
    entries = []
    for match in re.finditer(r"@\w+\{([^,]+),(.*?)(?=\n@|\Z)", content, flags=re.S):
        key = match.group(1).strip()
        body = match.group(2)
        fields = {"key": key}
        for field_match in re.finditer(r"\n\s*(\w+)\s*=\s*\{(.*?)\}\s*,?", body, flags=re.S):
            fields[field_match.group(1).lower()] = re.sub(r"\s+", " ", field_match.group(2)).strip()
        entries.append(fields)
    return entries


def build_user_guide_docx(path: Path) -> None:
    doc = Document()
    add_title(
        doc,
        "IVIVC App User Guide Review Copy",
        "Generated for Word comments/tracked changes. Apply accepted edits back to LaTeX source files.",
    )
    doc.add_heading("Reviewer note", level=1)
    add_paragraphs(
        doc,
        "This document is converted from the LaTeX user-guide source to support non-technical review in Word. "
        "Some equations, citations, and tables may be simplified for review convenience. The PDF generated from LaTeX remains the publication artifact."
    )

    user_guide_path = DOCS_DIR / "user_guide.tex"
    appendix_path = DOCS_DIR / "appendices" / "interpreting_model_performance.tex"
    if user_guide_path.exists():
        render_latex_file(doc, user_guide_path)
    if appendix_path.exists():
        doc.add_page_break()
        render_latex_file(doc, appendix_path, heading_prefix="Appendix")

    references = parse_bib_entries(DOCS_DIR / "references.bib")
    if references:
        doc.add_page_break()
        doc.add_heading("References", level=1)
        for entry in references:
            title = entry.get("title", "").replace("{", "").replace("}", "")
            author = entry.get("author", "")
            year = entry.get("year", "")
            citation = f"{entry['key']}: {author}. {title}. {year}."
            if entry.get("doi"):
                citation += f" DOI: {entry['doi']}."
            if entry.get("url"):
                citation += f" URL: {entry['url']}."
            doc.add_paragraph(citation, style="List Bullet")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def write_review_readme(path: Path) -> None:
    path.write_text(
        "# Review packet\n\n"
        "Generated files:\n\n"
        "- `app_content_review.docx`: app wording, modal/help text, model/metric/preprocessing summaries, and recurring result/report text.\n"
        "- `user_guide_review.docx`: Word-oriented review copy of the LaTeX user guide and appendix.\n\n"
        "These files are review artifacts only. Reviewers can use Word comments or tracked changes, but accepted edits should be applied back to source files such as `utilities/descriptions.yaml`, `docs/user_guide.tex`, `docs/appendices/interpreting_model_performance.tex`, and `docs/references.bib`.\n\n"
        "Screenshots or live-app review should be used separately for layout and visual feedback.\n",
        encoding="utf-8",
    )


def regenerate_user_guide() -> None:
    generator = DOCS_DIR / "generate_user_guide.py"
    if generator.exists():
        subprocess.run([sys.executable, str(generator)], cwd=DOCS_DIR, check=True)


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate Word review packet artifacts.")
    parser.add_argument("--output-dir", default=str(REVIEW_DIR), help="Directory for generated review files.")
    parser.add_argument("--skip-guide-regeneration", action="store_true", help="Use existing docs/user_guide.tex without regenerating it first.")
    args = parser.parse_args()

    output_dir = Path(args.output_dir)
    if not output_dir.is_absolute():
        output_dir = ROOT / output_dir
    output_dir.mkdir(parents=True, exist_ok=True)

    if not args.skip_guide_regeneration:
        regenerate_user_guide()

    app_docx = output_dir / "app_content_review.docx"
    guide_docx = output_dir / "user_guide_review.docx"
    build_app_content_docx(app_docx)
    build_user_guide_docx(guide_docx)
    write_review_readme(output_dir / "README.md")

    print(f"Wrote {app_docx.relative_to(ROOT)}")
    print(f"Wrote {guide_docx.relative_to(ROOT)}")
    print(f"Wrote {(output_dir / 'README.md').relative_to(ROOT)}")


if __name__ == "__main__":
    main()
