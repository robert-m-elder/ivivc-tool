"""Regression checks for faithful numerical and mathematical Word export."""

import ast
import base64
import io
from pathlib import Path
import unittest
import zipfile

from docx import Document
from docx.oxml.ns import nsmap
from lxml import etree

from utilities.word_report import _latex_to_omml, build_word_report


ROOT = Path(__file__).resolve().parents[1]


class WordReportTests(unittest.TestCase):
    def test_all_registered_equations_are_native_editable_math(self):
        equations = []
        for path in (ROOT / 'models').glob('*.py'):
            for node in ast.parse(path.read_text()).body:
                if isinstance(node, ast.Assign) and any(
                    isinstance(target, ast.Name) and target.id == 'latex_equation'
                    for target in node.targets
                ):
                    equations.append((path.name, ast.literal_eval(node.value)))
        self.assertGreaterEqual(len(equations), 16)
        for name, expression in equations:
            with self.subTest(model=name):
                math = _latex_to_omml(expression)
                self.assertEqual(etree.QName(math).localname, 'oMath')
                self.assertTrue(math.xpath('.//m:t', namespaces=nsmap))

    def test_grouping_nested_exponents_subscripts_and_fractions_survive(self):
        maximum = _latex_to_omml(r'y = \max\{1-kx-(x/t_f)^n,0\}')
        visible = ''.join(maximum.xpath('.//m:t/text()', namespaces=nsmap))
        self.assertIn('max{', visible)
        self.assertTrue(visible.endswith(',0}'))
        self.assertEqual(len(maximum.xpath('.//m:sSub', namespaces=nsmap)), 1)
        self.assertEqual(len(maximum.xpath('.//m:sSup', namespaces=nsmap)), 1)
        gompertz = _latex_to_omml(r'y = 1 - \exp\{-\exp[e r_{max}(t_{lag}-x)+1]\}')
        visible = ''.join(gompertz.xpath('.//m:t/text()', namespaces=nsmap))
        self.assertIn('exp{-exp[', visible)
        self.assertTrue(visible.endswith(']}'))
        nested = _latex_to_omml(r'y = a e^{b x^{c}}')
        self.assertEqual(len(nested.xpath('.//m:sSup', namespaces=nsmap)), 2)
        fraction = _latex_to_omml(r'y = \frac{a}{1+e^{-x}}')
        self.assertEqual(len(fraction.xpath('.//m:f/m:den/m:sSup', namespaces=nsmap)), 1)

    def test_complete_tables_unknown_equation_and_export_status(self):
        rows = [['Observation', 'Time', 'Response']]
        rows.extend([[index, index / 10, '0' if index == 0 else '1.23456e-08'] for index in range(37)])
        payload = {
            'report_label': 'Approach 3 Direct mapping',
            'analysis_timestamp': '2026-10-23 18:15 UTC',
            'image_export_summary': {'total': 2, 'exported': 1},
            'sections': [
                {'type': 'heading', 'level': 1, 'text': 'Numerical appendix'},
                {'type': 'table', 'rows': rows, 'title': 'All source observations'},
                {'type': 'equation', 'latex': r'y=\unknown{x+1}'},
                {'type': 'image', 'title': 'Missing figure', 'data_url': ''},
            ],
        }
        document = Document(build_word_report(payload))
        self.assertEqual(len(document.tables[0].rows), 38)
        self.assertEqual(document.tables[0].cell(1, 0).text, '0')
        self.assertEqual(document.tables[0].cell(1, 2).text, '0')
        self.assertEqual(document.tables[0].cell(37, 2).text, '1.23456e-08')
        paragraphs = '\n'.join(p.text for p in document.paragraphs)
        self.assertIn('Report export incomplete', paragraphs)
        self.assertIn('1 of 2 plot images', paragraphs)
        self.assertIn('Analysis recorded: 2026-10-23 18:15 UTC', paragraphs)
        self.assertIn('Exported:', paragraphs)
        self.assertIn('Equation source: y=\\unknown{x+1}', paragraphs)
        self.assertIn('Missing figure: image was unavailable', paragraphs)
        self.assertTrue(document.tables[0]._tbl.xpath('.//w:tr[1]/w:trPr/w:tblHeader'))
        self.assertIn('All source observations', paragraphs)

    def test_images_have_visible_captions_and_accessible_descriptions(self):
        # Tiny fixed PNG; no external images or plotting dependencies are needed.
        data = ('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8A'
                'AwMCAO+aK1cAAAAASUVORK5CYII=')
        stream = build_word_report({'sections': [
            {'type': 'image', 'title': 'Fitted curve', 'description': 'Shading shows the prediction band.',
             'data_url': 'data:image/png;base64,' + data},
        ]})
        document = Document(stream)
        self.assertEqual(len(document.inline_shapes), 1)
        captions = [p.text for p in document.paragraphs if p.style.name == 'Caption']
        self.assertEqual(captions, ['Fitted curve. Shading shows the prediction band.'])
        self.assertEqual(document.inline_shapes[0]._inline.docPr.get('descr'),
                         'Shading shows the prediction band.')
        self.assertLessEqual(document.inline_shapes[0].width,
                             document.sections[0].page_width - document.sections[0].left_margin
                             - document.sections[0].right_margin)


if __name__ == '__main__':
    unittest.main()
