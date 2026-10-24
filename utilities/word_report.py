"""Render structured analysis records as editable, accessible Word reports."""

import base64
import io
import re
from datetime import datetime, timezone

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.shared import RGBColor

WORD_MIME_TYPE = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
REPORT_LANGUAGE = 'en-US'




def _set_style_language(style, language=REPORT_LANGUAGE):
    """Set the proofing and screen-reader language for a Word style."""
    r_pr = style.element.get_or_add_rPr()
    language_node = r_pr.find(qn('w:lang'))
    if language_node is None:
        language_node = OxmlElement('w:lang')
        r_pr.append(language_node)
    language_node.set(qn('w:val'), language)
    language_node.set(qn('w:eastAsia'), language)
    language_node.set(qn('w:bidi'), language)


def _set_table_metadata(table, title, description, header_rows=1):
    """Add Word table metadata and identify repeated column-header rows."""
    tbl_pr = table._tbl.tblPr

    if title:
        caption = tbl_pr.find(qn('w:tblCaption'))
        if caption is None:
            caption = OxmlElement('w:tblCaption')
            tbl_pr.append(caption)
        caption.set(qn('w:val'), _clean_text(title))

    if description:
        desc = tbl_pr.find(qn('w:tblDescription'))
        if desc is None:
            desc = OxmlElement('w:tblDescription')
            tbl_pr.append(desc)
        desc.set(qn('w:val'), _clean_text(description))

    header_rows = max(0, min(int(header_rows or 0), len(table.rows)))
    for row in table.rows[:header_rows]:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn('w:tblHeader')) is None:
            tr_pr.append(OxmlElement('w:tblHeader'))


def _set_image_alt_text(inline_shape, title, description):
    """Set title and alternative text on an inline Word image."""
    doc_pr = inline_shape._inline.docPr
    if title:
        doc_pr.set('title', _clean_text(title))
    doc_pr.set('descr', _clean_text(description) or _clean_text(title) or 'IVIVC report plot')


def safe_report_filename(value):
    """Return a conservative report filename for download."""
    value = str(value or 'IVIVC Report')
    value = re.sub(r'[^A-Za-z0-9._ -]+', '', value).strip()
    value = re.sub(r'\s+', '_', value)
    return value[:80] or 'IVIVC_Report'


def _clean_text(value):
    """Normalize browser-extracted text before writing to Word."""
    value = '' if value is None else str(value)
    value = value.replace('\xa0', ' ')
    value = re.sub(r'\s+', ' ', value).strip()
    return value


def _clean_plot_alt_text(value):
    """Remove browser-only instructions from plot text used in Word alt text."""
    value = _clean_text(value)
    browser_only_sentences = (
        r'Use the Plotly toolbar to inspect or rescale the chart\.?',
        r'The controls beneath the plot can copy or download the plotted series '
        r'and any available error-bar bounds\.?',
    )
    for sentence in browser_only_sentences:
        value = re.sub(sentence, '', value, flags=re.IGNORECASE)
    return re.sub(r'\s+', ' ', value).strip()


def _add_document_styles(document):
    """Apply conservative styles and a document language to the report."""
    styles = document.styles
    styles['Normal'].font.name = 'Arial'
    styles['Normal'].font.size = Pt(10)
    styles['Normal'].paragraph_format.space_after = Pt(6)

    for style_name in ('Normal', 'Title', 'Subtitle', 'Caption', 'List Bullet'):
        if style_name in styles:
            _set_style_language(styles[style_name])
            styles[style_name].font.color.rgb = RGBColor(0, 0, 0)

    styles['Title'].font.name = 'Arial'
    styles['Title'].font.size = Pt(22)
    # Some python-docx templates carry a themed rule beneath the title.
    # Keep the report title plain in Word and in headless rendering.
    title_properties = styles['Title'].element.find(qn('w:pPr'))
    if title_properties is not None:
        for border in list(title_properties.findall(qn('w:pBdr'))):
            title_properties.remove(border)
    styles['Subtitle'].font.size = Pt(11)
    styles['Caption'].font.size = Pt(9)

    for style_name, size in [('Heading 1', 16), ('Heading 2', 13), ('Heading 3', 11)]:
        if style_name in styles:
            styles[style_name].font.name = 'Arial'
            styles[style_name].font.size = Pt(size)
            _set_style_language(styles[style_name])
            styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
            styles[style_name].paragraph_format.keep_with_next = True


def _usable_width(document):
    section = document.sections[-1]
    return section.page_width - section.left_margin - section.right_margin


def _table_column_widths(rows, total_width):
    """Allocate width to labels and long values without squeezing short numbers."""
    max_cols = max(len(row) for row in rows)
    weights = []
    for col in range(max_cols):
        values = [row[col] for row in rows if col < len(row)]
        lengths = sorted(len(value) for value in values)
        # A single verbose cell should wrap rather than consume the whole page.
        typical = lengths[min(len(lengths) - 1, int(.8 * len(lengths)))]
        header = len(values[0]) if values else 0
        weights.append(max(9, min(38, max(header, typical))))
    return [int(total_width * weight / sum(weights)) for weight in weights]


def _set_cell_fill(cell, color):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn('w:shd'))
    if shading is None:
        shading = OxmlElement('w:shd')
        properties.append(shading)
    shading.set(qn('w:fill'), color)


def _set_table_borders(table):
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = OxmlElement(f'w:{edge}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:color'), 'D9D9D9')
        borders.append(border)
    current = table._tbl.tblPr.find(qn('w:tblBorders'))
    if current is not None:
        table._tbl.tblPr.remove(current)
    table._tbl.tblPr.append(borders)


def _add_table(document, rows, title=None, description=None, header_rows=1):
    """Add an editable table with accessible title and header-row metadata."""
    rows = rows or []
    rows = [[_clean_text(cell) for cell in row] for row in rows]
    rows = [row for row in rows if any(row)]
    if not rows:
        return

    table_title = _clean_text(title) or 'Report data table'
    caption = document.add_paragraph(table_title, style='Caption')
    caption.paragraph_format.keep_with_next = True
    max_cols = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=max_cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = _table_column_widths(rows, _usable_width(document))
    for column, width in zip(table.columns, widths):
        column.width = width
    _set_table_borders(table)
    header_rows = max(0, min(int(header_rows or 0), len(rows)))

    for r_idx, row in enumerate(rows):
        for c_idx in range(max_cols):
            text = row[c_idx] if c_idx < len(row) else ''
            cell = table.cell(r_idx, c_idx)
            cell.width = widths[c_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margin = OxmlElement('w:tcMar')
            for side, value in (('top', 65), ('bottom', 65), ('left', 85), ('right', 85)):
                element = OxmlElement(f'w:{side}')
                element.set(qn('w:w'), str(value))
                element.set(qn('w:type'), 'dxa')
                margin.append(element)
            cell._tc.get_or_add_tcPr().append(margin)
            paragraph = cell.paragraphs[0]
            paragraph.text = text
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            paragraph.paragraph_format.keep_with_next = False
            if r_idx >= header_rows and re.fullmatch(r'[+\-\d.,eE% ±/]+|N/A', text):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for run in paragraph.runs:
                run.font.size = Pt(9)
            if r_idx < header_rows:
                _set_cell_fill(cell, 'DCE6F1')
                for run in paragraph.runs:
                    run.bold = True

    table_description = _clean_text(description) or f'{table_title}. Column headings are identified in the first row.'
    _set_table_metadata(table, table_title, table_description, header_rows=header_rows)
    document.add_paragraph()


def _decode_image_data_url(data_url):
    """Decode a browser data URL for insertion into a Word document."""
    if not data_url or not isinstance(data_url, str):
        return None
    match = re.match(r'^data:image/(png|jpeg|jpg);base64,(.+)$', data_url, re.IGNORECASE | re.DOTALL)
    if not match:
        return None
    try:
        return io.BytesIO(base64.b64decode(match.group(2)))
    except Exception:
        return None


def _add_image(document, data_url, title=None, description=None, caption=None):
    """Add a plot with a visible caption, accessible description and page-safe size."""
    image_stream = _decode_image_data_url(data_url)
    plot_title = _clean_text(title) or 'Report plot'
    plot_description = _clean_plot_alt_text(description) or plot_title

    if image_stream is None:
        document.add_paragraph(f'{plot_title}: image was unavailable for export.')
        return

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    try:
        inline_shape = run.add_picture(image_stream, width=_usable_width(document))
    except Exception:
        paragraph._element.getparent().remove(paragraph._element)
        document.add_paragraph(f'{plot_title}: image was unavailable for export.')
        return
    section = document.sections[-1]
    max_height = section.page_height - section.top_margin - section.bottom_margin - Inches(1)
    if inline_shape.height > max_height:
        scale = max_height / inline_shape.height
        inline_shape.width = int(inline_shape.width * scale)
        inline_shape.height = int(max_height)
    _set_image_alt_text(inline_shape, plot_title, plot_description)
    caption_text = _clean_plot_alt_text(caption) if caption else plot_description
    if caption_text != plot_title:
        caption_text = f'{plot_title}. {caption_text}'
    caption_paragraph = document.add_paragraph(caption_text, style='Caption')
    caption_paragraph.paragraph_format.keep_with_next = False


class _UnsupportedEquation(ValueError):
    """The supported LaTeX subset must not silently discard mathematical syntax."""


class _LatexMathParser:
    """Parse the small LaTeX subset used by the model registry, without evaluation.

    The resulting nodes retain grouping, subscripts, nested superscripts and
    fractions. Escaped braces are visible delimiters; unescaped braces group an
    argument, as in LaTeX. Unsupported commands are handled by an explicit fallback.
    """

    _plain_commands = {'exp', 'max', 'min', 'log', 'ln', 'sin', 'cos', 'tan', 'sign', 'abs'}
    _symbols = {'tau': 'τ', 'sigma': 'σ', 'mu': 'μ', 'alpha': 'α', 'beta': 'β',
                'Delta': 'Δ', 'epsilon': 'ε', 'infty': '∞', 'pm': '±',
                'cdot': '·', 'times': '×', 'le': '≤', 'ge': '≥'}

    def __init__(self, text):
        self.text = str(text or '').strip()
        self.index = 0

    def parse(self, closing=None):
        nodes = []
        while self.index < len(self.text):
            char = self.text[self.index]
            if char == '}':
                if closing != '}':
                    raise _UnsupportedEquation('Unmatched closing brace')
                self.index += 1
                return nodes
            node = self.atom()
            if node is None:
                continue
            superscript, subscript = None, None
            while self.index < len(self.text) and self.text[self.index] in '^_':
                marker = self.text[self.index]
                self.index += 1
                argument = self.argument()
                if marker == '^':
                    if superscript is not None:
                        raise _UnsupportedEquation('Repeated superscript')
                    superscript = argument
                else:
                    if subscript is not None:
                        raise _UnsupportedEquation('Repeated subscript')
                    subscript = argument
            if superscript is not None or subscript is not None:
                node = ('script', [node], subscript, superscript)
            nodes.append(node)
        if closing:
            raise _UnsupportedEquation('Unclosed brace')
        return nodes

    def argument(self):
        while self.index < len(self.text) and self.text[self.index].isspace():
            self.index += 1
        if self.index >= len(self.text):
            raise _UnsupportedEquation('Missing equation argument')
        if self.text[self.index] == '{':
            self.index += 1
            return self.parse('}')
        node = self.atom()
        if node is None:
            raise _UnsupportedEquation('Missing equation argument')
        return [node]

    def atom(self):
        char = self.text[self.index]
        self.index += 1
        if char == '{':
            return ('group', self.parse('}'))
        if char in '^_':
            raise _UnsupportedEquation('Script with no base')
        if char != '\\':
            return ('text', char, not char.isalpha())
        if self.index >= len(self.text):
            raise _UnsupportedEquation('Trailing backslash')
        escaped = self.text[self.index]
        if not escaped.isalpha():
            self.index += 1
            if escaped in '{}[]()|':
                return ('text', escaped, True)
            if escaped in ',;! ':
                return ('text', ' ', True)
            raise _UnsupportedEquation('Unsupported escaped character')
        start = self.index
        while self.index < len(self.text) and self.text[self.index].isalpha():
            self.index += 1
        command = self.text[start:self.index]
        if command in ('left', 'right'):
            return None
        if command == 'frac':
            return ('fraction', self.argument(), self.argument())
        if command in ('mathrm', 'operatorname', 'text'):
            return ('upright', self.argument())
        if command in self._plain_commands:
            return ('text', command, True)
        if command in self._symbols:
            return ('text', self._symbols[command], True)
        raise _UnsupportedEquation(f'Unsupported equation command: {command}')


def _math_run(value, upright=False):
    element = OxmlElement('m:r')
    if upright:
        properties = OxmlElement('m:rPr')
        style = OxmlElement('m:sty')
        style.set(qn('m:val'), 'p')
        properties.append(style)
        element.append(properties)
    text = OxmlElement('m:t')
    text.set(qn('xml:space'), 'preserve')
    text.text = value
    element.append(text)
    return element


def _append_math(parent, nodes, upright=False):
    for node in nodes:
        kind = node[0]
        if kind == 'text':
            parent.append(_math_run(node[1], upright or node[2]))
        elif kind in ('group', 'upright'):
            _append_math(parent, node[1], upright or kind == 'upright')
        elif kind == 'fraction':
            fraction = OxmlElement('m:f')
            for tag, contents in (('m:num', node[1]), ('m:den', node[2])):
                child = OxmlElement(tag)
                _append_math(child, contents, upright)
                fraction.append(child)
            parent.append(fraction)
        elif kind == 'script':
            base, subscript, superscript = node[1:]
            tag = 'm:sSubSup' if subscript is not None and superscript is not None else (
                'm:sSub' if subscript is not None else 'm:sSup')
            script = OxmlElement(tag)
            for child_tag, contents in (('m:e', base), ('m:sub', subscript), ('m:sup', superscript)):
                if contents is None:
                    continue
                child = OxmlElement(child_tag)
                _append_math(child, contents, upright)
                script.append(child)
            parent.append(script)


def _latex_to_omml(latex):
    """Return native editable Word math, or raise for an unsupported expression."""
    nodes = _LatexMathParser(latex).parse()
    if not nodes:
        raise _UnsupportedEquation('Empty equation')
    equation = OxmlElement('m:oMath')
    _append_math(equation, nodes)
    return equation


def _add_equation(document, latex, fallback_text=None):
    """Use native Word math; retain explicit source when a command is unsupported."""
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    try:
        equation = _latex_to_omml(latex or fallback_text)
    except _UnsupportedEquation:
        # Never strip grouping or guess the meaning of an unknown command.
        # A supplied canonical text equation is preferable; otherwise preserve
        # the original mathematical source in an explicitly identified form.
        if fallback_text and fallback_text != latex:
            paragraph.add_run(_clean_text(fallback_text))
        else:
            paragraph.add_run('Equation source: ' + _clean_text(latex or fallback_text))
        return
    paragraph._p.append(equation)


def _add_page_numbers(document):
    paragraph = document.sections[0].footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.add_run('Page ')
    page = OxmlElement('w:fldSimple')
    page.set(qn('w:instr'), 'PAGE')
    paragraph._p.append(page)
    for run in paragraph.runs:
        run.font.size = Pt(8)


def build_word_report(payload):
    """Build a DOCX report from structured content posted by the browser.

    The browser remains responsible for deciding which selected report is active
    and exporting Plotly figures to image data URLs. This function intentionally
    avoids any model selection or ranking logic.
    """
    payload = payload or {}
    report_label = _clean_text(payload.get('report_label')) or 'Selected IVIVC model'
    sections = payload.get('sections') or []

    document = Document()
    core = document.core_properties
    core.title = f'IVIVC Model Report - {report_label}'
    core.subject = 'Selected IVIVC model configuration, results, diagnostics, and plots'
    core.author = 'IVIVC App'
    core.keywords = 'IVIVC, absorbable materials, model fitting, cross-validation'
    core.category = 'IVIVC model report'
    core.comments = 'Generated by the IVIVC App.'

    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    _add_document_styles(document)
    _add_page_numbers(document)

    title = document.add_heading('IVIVC Model Report', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = document.add_paragraph(style='Subtitle')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(report_label).bold = True

    generated = document.add_paragraph()
    generated.alignment = WD_ALIGN_PARAGRAPH.CENTER
    analysis_timestamp = _clean_text(payload.get('analysis_timestamp'))
    if analysis_timestamp:
        generated.add_run(f'Analysis recorded: {analysis_timestamp}\n')
    generated.add_run(f'Exported: {datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")}')

    image_summary = payload.get('image_export_summary') or {}
    try:
        image_total = max(0, int(image_summary.get('total', 0)))
        image_exported = max(0, min(image_total, int(image_summary.get('exported', 0))))
    except (TypeError, ValueError):
        image_total = image_exported = 0
    if image_exported < image_total:
        incomplete = document.add_paragraph()
        incomplete.add_run('Report export incomplete. ').bold = True
        incomplete.add_run(
            f'{image_total-image_exported} of {image_total} plot images could not be included. '
            'The affected figures are identified below.'
        )

    document.add_paragraph()

    for item in sections:
        item_type = item.get('type')
        if item_type == 'heading':
            text = _clean_text(item.get('text'))
            if text:
                level = int(item.get('level') or 1)
                level = max(1, min(level, 3))
                document.add_heading(text, level=level)
        elif item_type == 'paragraph':
            text = _clean_text(item.get('text'))
            if text:
                document.add_paragraph(text)
        elif item_type == 'equation':
            _add_equation(document, item.get('latex'), item.get('text'))
        elif item_type == 'bullet':
            text = _clean_text(item.get('text'))
            if text:
                document.add_paragraph(text, style='List Bullet')
        elif item_type == 'table':
            _add_table(
                document,
                item.get('rows'),
                title=item.get('title'),
                description=item.get('description'),
                header_rows=item.get('header_rows', 1),
            )
        elif item_type == 'image':
            _add_image(
                document,
                item.get('data_url'),
                title=item.get('title'),
                description=item.get('description'),
                caption=item.get('caption'),
            )

    document_io = io.BytesIO()
    document.save(document_io)
    document_io.seek(0)
    return document_io
